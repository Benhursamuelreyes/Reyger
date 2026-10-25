"""Tests de la copia automática diaria, el aviso de stock bajo y el informe
de ventas por periodo."""

import os
import shutil
import sqlite3
import sys
import tkinter as tk
from datetime import date, datetime, timedelta
from types import SimpleNamespace

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

import reyger.core.db as modulo_db
from reyger.core import backup, migrations
from reyger.domain import informes
from reyger.domain import inventario as dominio_inventario
from reyger.domain import ventas as dv

PLANTILLA = os.path.join(
    os.path.dirname(__file__), "..", "src", "reyger", "assets", "database.db"
)


@pytest.fixture
def bd(tmp_path, monkeypatch):
    ruta = str(tmp_path / "database.db")
    shutil.copyfile(PLANTILLA, ruta)
    conn = sqlite3.connect(ruta)
    migrations.run_migrations(conn)
    conn.executemany(
        "INSERT INTO inventario (nombre, proveedor, precio, costo, stock, tipo_iva)"
        " VALUES (?, 'Prov', ?, 1.0, ?, ?)",
        [("Café", 2.20, 50, 10), ("Libro", 12.10, 2, 4), ("Pan", 1.00, 0, 4)],
    )
    conn.commit()
    conn.close()
    monkeypatch.setattr(modulo_db, "get_db_path", lambda: ruta)
    modulo_db.close()
    yield ruta
    modulo_db.close()


# ------------------------------------------------------- copia automática
def _filas_inventario(ruta):
    conn = sqlite3.connect(ruta)
    try:
        return conn.execute("SELECT COUNT(*) FROM inventario").fetchone()[0]
    finally:
        conn.close()


def test_respaldo_automatico_crea_copia_completa(bd):
    ruta = backup.respaldo_automatico(bd, hoy=datetime(2026, 10, 25))
    assert ruta.endswith(os.path.join("respaldos", "automatico_20261025.db"))
    assert _filas_inventario(ruta) == 3
    assert not os.path.exists(ruta + ".tmp")


def test_respaldo_automatico_sobrescribe_el_del_dia(bd):
    hoy = datetime(2026, 10, 25)
    backup.respaldo_automatico(bd, hoy=hoy)
    modulo_db.execute(
        "INSERT INTO inventario (nombre, proveedor, precio, costo, stock) VALUES ('Té','',1,1,1)"
    )
    ruta = backup.respaldo_automatico(bd, hoy=hoy)
    assert _filas_inventario(ruta) == 4  # incluye lo vendido/creado después
    assert len(os.listdir(os.path.dirname(ruta))) == 1


def test_respaldo_automatico_conserva_siete_dias(bd):
    inicio = datetime(2026, 10, 1)
    for dia in range(10):
        ruta = backup.respaldo_automatico(bd, hoy=inicio + timedelta(days=dia))
    copias = sorted(os.listdir(os.path.dirname(ruta)))
    assert len(copias) == backup.DIAS_RESPALDO_AUTOMATICO
    assert copias[0] == "automatico_20261004.db" and copias[-1] == "automatico_20261010.db"


def test_respaldo_automatico_sin_base(tmp_path):
    assert backup.respaldo_automatico(str(tmp_path / "no_existe.db")) is None


def test_respaldo_no_mezcla_con_los_previos_a_importar(bd):
    for _ in range(backup.MAX_RESPALDOS + 2):
        backup.respaldar_bd_actual(bd)
    ruta = backup.respaldo_automatico(bd)
    assert os.path.exists(ruta)  # la purga de los .bak no la toca


def test_cerrar_hace_la_copia_y_cierra(monkeypatch):
    from reyger.ui.manager import Manager

    llamadas = []
    monkeypatch.setattr(backup, "respaldo_automatico", lambda: llamadas.append("copia"))
    ventana = SimpleNamespace(destroy=lambda: llamadas.append("cerrar"))
    Manager._al_cerrar(ventana)
    assert llamadas == ["copia", "cerrar"]


def test_fallo_en_la_copia_no_impide_cerrar(monkeypatch, caplog):
    from reyger.ui.manager import Manager

    def romper():
        raise OSError("disco lleno")

    monkeypatch.setattr(backup, "respaldo_automatico", romper)
    cerrada = []
    Manager._al_cerrar(SimpleNamespace(destroy=lambda: cerrada.append(True)))
    assert cerrada == [True]
    assert "disco lleno" in caplog.text


def test_copia_desactivada_en_ajustes(monkeypatch):
    from reyger.config import ConfigManager
    from reyger.ui.manager import Manager

    ConfigManager().set("respaldo_automatico", False)
    monkeypatch.setattr(
        backup, "respaldo_automatico", lambda: pytest.fail("no debía copiar")
    )
    cerrada = []
    Manager._al_cerrar(SimpleNamespace(destroy=lambda: cerrada.append(True)))
    assert cerrada == [True]


# ------------------------------------------------------------- stock bajo
@pytest.mark.parametrize("stock, esperado", [
    (0, "agotado"), (-1, "agotado"), (1, "bajo"), (5, "bajo"), (6, None), ("x", None),
])
def test_estado_stock(stock, esperado):
    assert dominio_inventario.estado_stock(stock, 5) == esperado


def test_inventario_marca_y_cuenta_el_stock_bajo(bd):
    from reyger.ui.inventario import Inventario

    try:
        raiz = tk.Tk()
    except tk.TclError:
        pytest.skip("sin display para Tk")
    raiz.withdraw()
    try:
        inv = Inventario(raiz)
        etiquetas = {
            inv.tre.item(i, "values")[1]: inv.tre.item(i, "tags")
            for i in inv.tre.get_children()
        }
        assert etiquetas["Café"] in ("", ())
        assert "bajo" in etiquetas["Libro"]
        assert "agotado" in etiquetas["Pan"]
        texto = inv.lbl_stock_bajo.cget("text")
        assert "1 agotado" in texto and "1 con 5 unidades o menos" in texto
    finally:
        raiz.destroy()


# ---------------------------------------------------------------- informes
def _linea(nombre, cantidad):
    precios = {"Café": (2.20, 10.0), "Libro": (12.10, 4.0)}
    precio, iva = precios[nombre]
    return dv.LineaVenta(nombre, precio, cantidad, iva)


def _fechar(numero, fecha):
    """Fija la fecha (hora local ISO) del documento *numero*."""
    modulo_db.execute("UPDATE documentos SET fecha = ? WHERE numero = ?", (fecha, numero))


def test_informe_suma_solo_lo_emitido_en_el_periodo(bd):
    t1 = dv.registrar_venta([_linea("Café", 2), _linea("Libro", 1)], "ticket", "Efectivo", 20)
    t2 = dv.registrar_venta([_linea("Café", 1)], "ticket", "Tarjeta", 0, 5)
    f1 = dv.registrar_venta([_linea("Libro", 1)], "factura", "Tarjeta", 0, 20)
    fuera = dv.registrar_venta([_linea("Café", 5)], "ticket", "Efectivo", 20)
    dv.guardar_borrador([_linea("Libro", 3)], None, "05/10/2026 10:00")  # no cuenta

    _fechar(t1.numero, "2026-10-05 12:00:00")
    _fechar(t2.numero, "2026-10-06 09:15:00")
    _fechar(f1.numero, "2026-10-06 23:59:59")  # último segundo del día
    _fechar(fuera.numero, "2026-11-01 00:00:00")

    inf = informes.informe_ventas(date(2026, 10, 1), date(2026, 10, 31))

    assert inf.documentos == 3
    assert inf.total == pytest.approx(4.40 + 12.10 + 2.20 + 12.10)
    assert inf.base + inf.cuota == pytest.approx(inf.total, abs=0.02)
    assert inf.por_dia == [
        ("2026-10-05", 1, pytest.approx(16.50)),
        ("2026-10-06", 2, pytest.approx(14.30)),
    ]
    assert dict((t, d) for t, d, _ in inf.por_tipo) == {"factura": 1, "ticket": 2}
    assert dict((m, d) for m, d, _ in inf.por_metodo) == {"Efectivo": 1, "Tarjeta": 2}
    assert inf.por_producto[0][:2] == ("Libro", 2)
    assert inf.por_producto[1][:2] == ("Café", 3)


def test_informe_periodo_vacio_y_fechas_invertidas(bd):
    inf = informes.informe_ventas(date(2020, 1, 1), date(2020, 1, 31))
    assert inf.documentos == 0 and inf.total == 0.0 and inf.por_dia == []
    with pytest.raises(ValueError):
        informes.informe_ventas(date(2026, 2, 1), date(2026, 1, 1))


def test_factura_emitida_desde_borrador_lleva_la_fecha_de_emision(bd):
    borrador = dv.guardar_borrador([_linea("Café", 1)], None, "01/01/2020 09:00")
    doc = dv.emitir_borrador(borrador.numero)
    hoy = date.today()
    inf = informes.informe_ventas(hoy - timedelta(days=1), hoy + timedelta(days=1))
    assert inf.documentos == 1
    assert modulo_db.query_one(
        "SELECT date(fecha) AS dia FROM documentos WHERE numero = ?", (doc.numero,)
    )["dia"] == hoy.isoformat()


def test_exportar_informe_a_excel(bd, tmp_path):
    from openpyxl import load_workbook

    doc = dv.registrar_venta([_linea("Café", 2)], "ticket", "Efectivo", 10)
    _fechar(doc.numero, "2026-10-05 12:00:00")
    inf = informes.informe_ventas(date(2026, 10, 1), date(2026, 10, 31))
    ruta = informes.exportar_excel(inf, str(tmp_path / "informe.xlsx"))

    libro = load_workbook(ruta)
    assert libro.sheetnames == ["Resumen", "Por día", "Por producto", "Por método de pago"]
    resumen = {fila[0]: fila[1] for fila in libro["Resumen"].iter_rows(min_row=2, values_only=True)}
    assert resumen["Documentos emitidos"] == 1
    assert resumen["Total"] == pytest.approx(4.40)
    dia = list(libro["Por día"].iter_rows(min_row=2, values_only=True))[0]
    assert dia[0].date() == date(2026, 10, 5)


# ----------------------------------------------------------- integridad
def test_copia_danada_no_sustituye_a_la_buena(bd, monkeypatch):
    hoy = datetime(2026, 10, 25)
    buena = backup.respaldo_automatico(bd, hoy=hoy)
    contenido = open(buena, "rb").read()
    monkeypatch.setattr(backup, "comprobar_integridad", lambda ruta: ["página 7 corrupta"])
    with pytest.raises(backup.BackupError, match="página 7 corrupta"):
        backup.respaldo_automatico(bd, hoy=hoy)
    assert open(buena, "rb").read() == contenido
    assert not os.path.exists(buena + ".tmp")


def test_comprobar_integridad_detecta_claves_foraneas_rotas(bd):
    assert backup.comprobar_integridad(bd) == []
    conn = sqlite3.connect(bd)
    conn.execute("PRAGMA foreign_keys = OFF")
    conn.execute(
        "INSERT INTO documentos (numero, tipo, estado, cliente_id) VALUES ('T-1','ticket','emitido', 999)"
    )
    conn.commit()
    conn.close()
    problemas = backup.comprobar_integridad(bd)
    assert problemas and "documentos" in problemas[0]


def test_al_cerrar_avisa_si_la_base_esta_danada(monkeypatch):
    from reyger.ui import manager as modulo_manager

    def danada():
        raise backup.BackupError("errores de integridad")

    avisos, cerrada = [], []
    monkeypatch.setattr(backup, "respaldo_automatico", danada)
    monkeypatch.setattr(
        modulo_manager.messagebox, "showwarning", lambda *a, **k: avisos.append(a[1])
    )
    modulo_manager.Manager._al_cerrar(SimpleNamespace(destroy=lambda: cerrada.append(True)))
    assert avisos == ["errores de integridad"] and cerrada == [True]


def test_optimizar_sin_conexion_no_falla():
    modulo_db.close()
    modulo_db.optimizar()
