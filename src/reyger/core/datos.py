"""Importación/exportación granular de clientes y productos (catálogo).

Complementa a :mod:`reyger.core.backup`, que trabaja con la base de
datos completa. Aquí cada tipo de datos se gestiona por separado en
``.xlsx`` o ``.csv``:

- **Productos / catálogo** (tabla ``inventario``): se identifican por su
  Código/SKU (columna ``codigo_barras``).
- **Clientes** (tabla ``clientes``): se identifican por su NIF/CIF
  (columna ``documento``).

Los ficheros emparejan columnas por *nombre de cabecera* (con alias
habituales como "Precio de Coste" o "Precio de Venta"), de modo que el
orden de las columnas no importa. Frente a registros duplicados hay dos
modos:

- ``"actualizar"``: si el identificador (Código/SKU o NIF/CIF) ya
  existe, actualiza sus datos.
- ``"ignorar"``: deja el registro local intacto y solo añade los nuevos.

Cada importación valida el fichero completo antes de abrir la
transacción; si cualquier fila es inválida, **no se modifica nada** y se
informan las filas con errores. Antes de aplicar cambios se genera una
copia de seguridad automática de la base actual.
"""

import csv
import os
import sqlite3
import unicodedata

from .backup import EXCEL_DISPONIBLE, BackupError, respaldar_bd_actual
from ..domain.fiscal import TIPOS_IVA, normalizar_tipo_iva
from ..resources import get_db_path

# ---------------------------------------------------------------------------
# Claves canónicas: (clave_interna, etiqueta_de_plantilla)
# ---------------------------------------------------------------------------

CAMPOS_PRODUCTO = (
    ("codigo", "Código"),
    ("nombre", "Nombre"),
    ("costo", "Precio Coste"),
    ("precio", "Precio Venta"),
    ("stock", "Stock"),
    ("tipo_iva", "IVA"),
    ("proveedor", "Proveedor"),
    ("categoria", "Categoría"),
    ("margen_porcentaje", "Margen %"),
)

CAMPOS_CLIENTE = (
    ("nombre", "Nombre"),
    ("documento", "NIF/CIF"),
    ("telefono", "Teléfono"),
    ("email", "Email"),
    ("direccion", "Dirección"),
    ("codigo_postal", "Código Postal"),
    ("provincia", "Provincia"),
    ("notas", "Notas"),
)

#: Nombres de columna aceptados (normalizados sin acentos) por clave.
ALIASES_PRODUCTO = {
    "codigo": ("codigo", False),
    "sku": ("codigo", False),
    "codigo/sku": ("codigo", False),
    "codigo del producto": ("codigo", False),
    "codigo de barras": ("codigo", False),
    "codigo_barras": ("codigo", False),
    "barras": ("codigo", False),
    "nombre": ("nombre", True),
    "nombre producto": ("nombre", True),
    "producto": ("nombre", True),
    "articulo": ("nombre", True),
    "descripcion": ("nombre", True),
    "precio coste": ("costo", True),
    "precio de coste": ("costo", True),
    "costo": ("costo", True),
    "coste": ("costo", True),
    "coste unitario": ("costo", True),
    "precio venta": ("precio", True),
    "precio de venta": ("precio", True),
    "pvp": ("precio", True),
    "precio": ("precio", True),
    "venta": ("precio", True),
    "stock": ("stock", True),
    "existencias": ("stock", True),
    "unidades": ("stock", True),
    "cantidad": ("stock", True),
    "iva": ("tipo_iva", True),
    "tipo iva": ("tipo_iva", True),
    "tipo_iva": ("tipo_iva", True),
    "tipo de iva": ("tipo_iva", True),
    "iva (%)": ("tipo_iva", True),
    "iva %": ("tipo_iva", True),
    "proveedor": ("proveedor", False),
    "categoria": ("categoria", False),
    "categoria_id": ("categoria", False),
    "margen": ("margen_porcentaje", False),
    "margen %": ("margen_porcentaje", False),
    "margen (%)": ("margen_porcentaje", False),
    "margen_porcentaje": ("margen_porcentaje", False),
}

ALIASES_CLIENTE = {
    "nombre": ("nombre", True),
    "razon social": ("nombre", True),
    "nombre/razon social": ("nombre", True),
    "cliente": ("nombre", True),
    "nif/cif": ("documento", False),
    "nif": ("documento", False),
    "cif": ("documento", False),
    "nie": ("documento", False),
    "dni": ("documento", False),
    "documento": ("documento", False),
    "nif/cif/nie": ("documento", False),
    "telefono": ("telefono", False),
    "tel": ("telefono", False),
    "email": ("email", False),
    "e-mail": ("email", False),
    "correo": ("email", False),
    "correo electronico": ("email", False),
    "direccion": ("direccion", False),
    "domicilio": ("direccion", False),
    "codigo postal": ("codigo_postal", False),
    "cp": ("codigo_postal", False),
    "provincia": ("provincia", False),
    "notas": ("notas", False),
    "observaciones": ("notas", False),
    "tipo documento": ("tipo_documento", False),
    "tipo_documento": ("tipo_documento", False),
    "tipo doc": ("tipo_documento", False),
}


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------


def _normalizar(texto):
    """Minúsculas, sin acentos y con espacios colapsados para comparar."""
    texto = unicodedata.normalize("NFKD", str(texto or ""))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return " ".join(texto.lower().split())


def _vacio(valor):
    return valor is None or str(valor).strip() == ""


def _a_float(valor):
    if _vacio(valor):
        return None
    texto = str(valor).strip().replace(" ", "").replace("€", "")
    if not texto:
        return None
    try:
        return float(texto.replace(",", "."))
    except ValueError:
        return None


def _a_int(valor):
    numero = _a_float(valor)
    if numero is None:
        return None
    return int(numero)


def _a_iva(valor):
    """Tipo de IVA válido (21/10/4) o ``None`` si está vacío o es inválido."""
    if _vacio(valor):
        return None
    tipo = normalizar_tipo_iva(valor)
    if tipo is None or tipo not in TIPOS_IVA:
        return None
    return tipo


def _a_texto(valor, mayusculas=False):
    if _vacio(valor):
        return None
    texto = str(valor).strip()
    return texto.upper() if mayusculas else texto


def _detectar_tipo_documento(documento):
    """Infiere NIF/NIE/CIF de un documento español (heurística ligera)."""
    if not documento:
        return None
    primera = documento[0]
    if primera in "XYZ":
        return "NIE"
    if primera.isalpha():
        return "CIF"
    return "NIF"


# ---------------------------------------------------------------------------
# Lectura de ficheros (.xlsx / .csv)
# ---------------------------------------------------------------------------


def _leer_archivo(ruta):
    """Devuelve ``(cabecera, filas)`` del primer libro/CSV de ``ruta``.

    Descarta filas completamente vacías y convierte a None las celdas
    en blanco. Un libro Excel sin hojas o el fichero sin filas se
    rechazan con un mensaje descriptivo.
    """
    extension = os.path.splitext(str(ruta))[1].lower()
    filas = []
    if extension == ".xlsx":
        if not EXCEL_DISPONIBLE:
            raise BackupError(
                "El soporte de Excel requiere el paquete 'openpyxl', que no "
                "está instalado en este equipo."
            )
        try:
            from openpyxl import load_workbook

            libro = load_workbook(str(ruta), read_only=True, data_only=True)
        except Exception as e:
            raise BackupError(f"No se pudo abrir el ficho de Excel: {e}") from e
        try:
            if not libro.worksheets:
                raise BackupError(
                    "El libro de Excel no contiene ninguna hoja de cálculo."
                )
            hoja = libro.worksheets[0]
            for fila in hoja.iter_rows(values_only=True):
                valores = [None if _vacio(v) else v for v in (fila or ())]
                if any(v is not None for v in valores):
                    filas.append(valores)
        finally:
            libro.close()
    elif extension == ".csv":
        try:
            with open(str(ruta), newline="", encoding="utf-8-sig") as f:
                lector = csv.reader(f)
                for fila in lector:
                    valores = [None if _vacio(v) else v for v in fila]
                    if any(v is not None for v in valores):
                        filas.append(valores)
        except OSError as e:
            raise BackupError(f"No se pudo leer el fichero CSV: {e}") from e
    else:
        raise BackupError(
            "Formato no reconocido. Use un fichero .xlsx o .csv."
        )

    if not filas:
        raise BackupError("El fichero está vacío o no contiene datos.")
    cabecera = [str(c).strip() for c in filas[0]]
    if not cabecera or not any(cabecera):
        raise BackupError(
            "La primera fila debe contener la cabecera con los nombres "
            "de las columnas."
        )
    return cabecera, filas[1:]


def _resolver(cabecera, alias, campos):
    """Mapea la cabecera a claves internas.

    Devuelve ``(mapa, faltantes)`` donde ``mapa`` es ``{clave: indice}``
    y *faltantes* son las columnas obligatorias que no aparecen.
    """
    mapa = {}
    for indice, nombre in enumerate(cabecera):
        clave = alias.get(_normalizar(nombre))
        if clave is not None and clave[0] not in mapa:
            mapa[clave[0]] = indice
    requeridas = {
        clave for clave, _ in campos if clave in alias and alias[clave][1]
    }
    faltantes = sorted(requeridas - set(mapa))
    return mapa, faltantes


def _celda(fila, mapa, clave):
    indice = mapa.get(clave)
    if indice is None or indice >= len(fila):
        return None
    return fila[indice]


def _etiquetas(campos):
    return [etiqueta for _, etiqueta in campos]


def _escribir(ruta, cabecera, filas, nombre_hoja):
    """Escribe ``cabecera``+``filas`` en .xlsx o .csv según la extensión."""
    extension = os.path.splitext(str(ruta))[1].lower()
    if extension == ".xlsx":
        if not EXCEL_DISPONIBLE:
            raise BackupError(
                "El soporte de Excel requiere el paquete 'openpyxl', que no "
                "está instalado en este equipo."
            )
        from openpyxl import Workbook

        libro = Workbook(write_only=True)
        try:
            hoja = libro.create_sheet((nombre_hoja or "Datos")[:31])
            hoja.append(list(cabecera))
            for fila in filas:
                hoja.append(list(fila))
            libro.save(str(ruta))
        finally:
            libro.close()
    elif extension == ".csv":
        try:
            with open(str(ruta), "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.writer(f)
                writer.writerow(cabecera)
                for fila in filas:
                    writer.writerow(["" if v is None else v for v in fila])
        except OSError as e:
            raise BackupError(f"No se pudo escribir el fichero CSV: {e}") from e
    else:
        raise BackupError(
            "Formato no reconocido. Use una extensión .xlsx o .csv."
        )
    return ruta


# ---------------------------------------------------------------------------
# Plantillas de ejemplo
# ---------------------------------------------------------------------------


def plantilla_productos(ruta):
    """Genera una plantilla .xlsx/.csv de catálogo con un ejemplo."""
    ejemplo = [
        "1234567890", "Producto de ejemplo", 5.00, 7.50, 100, 21,
        "Proveedor de ejemplo", "General", 50,
    ]
    return _escribir(ruta, _etiquetas(CAMPOS_PRODUCTO), [ejemplo], "Productos")


def plantilla_clientes(ruta):
    """Genera una plantilla .xlsx/.csv de clientes con un ejemplo."""
    ejemplo = [
        "Cliente de ejemplo", "12345678A", "600 000 000",
        "cliente@ejemplo.es", "C/ Mayor 1, 28001 Madrid", "28001",
        "Madrid", "Cliente de prueba",
    ]
    return _escribir(ruta, _etiquetas(CAMPOS_CLIENTE), [ejemplo], "Clientes")


# ---------------------------------------------------------------------------
# Exportación
# ---------------------------------------------------------------------------


def exportar_productos(ruta, db_path=None):
    """Exporta el catálogo (tabla ``inventario``) a .xlsx o .csv."""
    conn = sqlite3.connect(str(db_path or get_db_path()))
    try:
        cursor = conn.execute(
            """
            SELECT i.codigo_barras, i.nombre, i.costo, i.precio, i.stock,
                   i.tipo_iva, i.proveedor, c.nombre,
                   i.margen_porcentaje
            FROM inventario i
            LEFT JOIN categorias c ON c.id = i.categoria_id
            ORDER BY i.nombre
            """
        )
        filas = [
            [
                fila[0], fila[1], fila[2], fila[3], fila[4], fila[5],
                fila[6], fila[7], fila[8],
            ]
            for fila in cursor.fetchall()
        ]
    finally:
        conn.close()
    return _escribir(ruta, _etiquetas(CAMPOS_PRODUCTO), filas, "Productos")


def exportar_clientes(ruta, db_path=None):
    """Exporta los clientes (tabla ``clientes``) a .xlsx o .csv."""
    conn = sqlite3.connect(str(db_path or get_db_path()))
    try:
        cursor = conn.execute(
            """
            SELECT nombre, documento, telefono, email, direccion,
                   codigo_postal, provincia, notas
            FROM clientes
            ORDER BY nombre
            """
        )
        filas = [list(fila) for fila in cursor.fetchall()]
    finally:
        conn.close()
    return _escribir(ruta, _etiquetas(CAMPOS_CLIENTE), filas, "Clientes")


# ---------------------------------------------------------------------------
# Importación
# ---------------------------------------------------------------------------


def _resumen(tipo, insertados, actualizados, ignorados, respaldo):
    return {
        "insertados": insertados,
        "actualizados": actualizados,
        "ignorados": ignorados,
        "respaldo": respaldo,
        "tipo": tipo,
    }


def importar_productos(ruta, modo="actualizar", db_path=None):
    """Importa el catálogo desde .xlsx/.csv (UPSERT por Código/SKU).

    ``modo`` puede ser ``"actualizar"`` (por defecto) o ``"ignorar"``.
    Devuelve un diccionario con el recuento de insertados, actualizados e
    ignorados. Antes de tocar nada valida todo el fichero: si hay
    cualquier error, lanza :class:`BackupError` sin modificar la base.
    """
    if modo not in ("actualizar", "ignorar"):
        raise BackupError(
            f"Modo de importación desconocido: «{modo}» (use 'actualizar' "
            "o 'ignorar')."
        )
    cabecera, filas = _leer_archivo(ruta)
    mapa, faltantes = _resolver(cabecera, ALIASES_PRODUCTO, CAMPOS_PRODUCTO)
    if faltantes:
        etiquetas = dict(CAMPOS_PRODUCTO)
        raise BackupError(
            "El fichero de productos no tiene las columnas obligatorias: "
            + ", ".join(etiquetas.get(clave, clave) for clave in faltantes)
            + ". Use la plantilla de ejemplo para conocer los nombres "
            "esperados."
        )

    registros = []
    errores = []
    for numero, fila in enumerate(filas, start=2):
        nombre = _a_texto(_celda(fila, mapa, "nombre"))
        costo = _a_float(_celda(fila, mapa, "costo"))
        precio = _a_float(_celda(fila, mapa, "precio"))
        stock = _a_int(_celda(fila, mapa, "stock"))
        comprobar = (
            (f"fila {numero}: falta el Nombre", not nombre),
            (f"fila {numero}: Precio de Coste inválido",
             costo is None),
            (f"fila {numero}: Precio de Venta inválido", precio is None),
            (f"fila {numero}: Stock inválido", stock is None),
        )
        for mensaje, mal in comprobar:
            if mal:
                errores.append(mensaje)
                break
        iva = _a_iva(_celda(fila, mapa, "tipo_iva"))
        if _celda(fila, mapa, "tipo_iva") is not None and iva is None:
            errores.append(
                f"fila {numero}: tipo de IVA inválido (use 21, 10 o 4)"
            )
        codigo = _a_texto(_celda(fila, mapa, "codigo"), mayusculas=False)
        proveedor = _a_texto(_celda(fila, mapa, "proveedor")) or ""
        categoria = _a_texto(_celda(fila, mapa, "categoria"))
        margen = _a_float(_celda(fila, mapa, "margen_porcentaje"))
        if margen is None and costo and precio and costo > 0:
            margen = round((precio - costo) / costo * 100, 2)
        registros.append({
            "codigo": codigo,
            "nombre": nombre,
            "costo": costo,
            "precio": precio,
            "stock": stock,
            "tipo_iva": iva if iva is not None else 21.0,
            "proveedor": proveedor,
            "proveedor_id": None,
            "categoria": categoria,
            "categoria_id": None,
            "margen_porcentaje": margen,
        })

    if errores:
        _lanzar_errores("productos", errores)

    destino = str(db_path or get_db_path())
    respaldo = respaldar_bd_actual(destino)
    conn = sqlite3.connect(destino)
    insertados = actualizados = ignorados = 0
    try:
        for registro in registros:
            if registro["codigo"]:
                fila_existente = conn.execute(
                    "SELECT id FROM inventario WHERE codigo_barras = ?",
                    (registro["codigo"],),
                ).fetchone()
            else:
                fila_existente = None

            if registro["proveedor"]:
                fila_prov = conn.execute(
                    "SELECT id FROM proveedores WHERE nombre = ?",
                    (registro["proveedor"],),
                ).fetchone()
                registro["proveedor_id"] = fila_prov[0] if fila_prov else None
            if registro["categoria"]:
                fila_cat = conn.execute(
                    "SELECT id FROM categorias WHERE nombre = ?",
                    (registro["categoria"],),
                ).fetchone()
                registro["categoria_id"] = fila_cat[0] if fila_cat else None

            if fila_existente is not None:
                if modo == "ignorar":
                    ignorados += 1
                    continue
                conn.execute(
                    """
                    UPDATE inventario
                    SET nombre=?, proveedor=?, precio=?, costo=?, stock=?,
                        proveedor_id=?, tipo_iva=?, categoria_id=?, codigo_barras=?, 
                        margen_porcentaje=?
                    WHERE id=?
                    """,
                    (
                        registro["nombre"], registro["proveedor"],
                        registro["precio"], registro["costo"],
                        registro["stock"], registro["proveedor_id"],
                        registro["tipo_iva"], registro["categoria_id"],
                        registro["codigo"], registro["margen_porcentaje"],
                        fila_existente[0],
                    ),
                )
                actualizados += 1
            else:
                conn.execute(
                    """
                    INSERT INTO inventario (nombre, proveedor, precio, costo,
                        stock, proveedor_id, tipo_iva, categoria_id, codigo_barras,
                        margen_porcentaje)
                    VALUES (?,?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        registro["nombre"], registro["proveedor"],
                        registro["precio"], registro["costo"],
                        registro["stock"], registro["proveedor_id"],
                        registro["tipo_iva"], registro["categoria_id"],
                        registro["codigo"], registro["margen_porcentaje"],
                    ),
                )
                insertados += 1
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        raise BackupError(f"Error importando productos: {e}") from e
    finally:
        conn.close()
    return _resumen("productos", insertados, actualizados, ignorados, respaldo)


def importar_clientes(ruta, modo="actualizar", db_path=None):
    """Importa clientes desde .xlsx/.csv (UPSERT por NIF/CIF).

    ``modo`` puede ser ``"actualizar"`` (por defecto) o ``"ignorar"``.
    Devuelve un diccionario con el recuento de insertados, actualizados e
    ignorados. Antes de tocar nada valida todo el fichero: si hay
    cualquier error, lanza :class:`BackupError` sin modificar la base.
    """
    if modo not in ("actualizar", "ignorar"):
        raise BackupError(
            f"Modo de importación desconocido: «{modo}» (use 'actualizar' "
            "o 'ignorar')."
        )
    cabecera, filas = _leer_archivo(ruta)
    mapa, faltantes = _resolver(cabecera, ALIASES_CLIENTE, CAMPOS_CLIENTE)
    if faltantes:
        etiquetas = dict(CAMPOS_CLIENTE)
        raise BackupError(
            "El fichero de clientes no tiene las columnas obligatorias: "
            + ", ".join(etiquetas.get(clave, clave) for clave in faltantes)
            + ". Use la plantilla de ejemplo para conocer los nombres "
            "esperados."
        )

    registros = []
    errores = []
    for numero, fila in enumerate(filas, start=2):
        nombre = _a_texto(_celda(fila, mapa, "nombre"))
        documento = _a_texto(_celda(fila, mapa, "documento"), mayusculas=True)
        tipo_documento = _a_texto(_celda(fila, mapa, "tipo_documento"))
        if not nombre:
            errores.append(f"fila {numero}: falta el Nombre del cliente")
        if tipo_documento is None:
            tipo_documento = _detectar_tipo_documento(documento) or "NIF"
        registros.append({
            "nombre": nombre,
            "documento": documento,
            "tipo_documento": tipo_documento,
            "telefono": _a_texto(_celda(fila, mapa, "telefono")),
            "email": _a_texto(_celda(fila, mapa, "email")),
            "direccion": _a_texto(_celda(fila, mapa, "direccion")),
            "codigo_postal": _a_texto(_celda(fila, mapa, "codigo_postal")),
            "provincia": _a_texto(_celda(fila, mapa, "provincia")),
            "notas": _a_texto(_celda(fila, mapa, "notas")),
        })

    if errores:
        _lanzar_errores("clientes", errores)

    destino = str(db_path or get_db_path())
    respaldo = respaldar_bd_actual(destino)
    conn = sqlite3.connect(destino)
    insertados = actualizados = ignorados = 0
    try:
        for registro in registros:
            if registro["documento"]:
                fila_existente = conn.execute(
                    "SELECT id FROM clientes WHERE documento = ?",
                    (registro["documento"],),
                ).fetchone()
            else:
                fila_existente = None

            if fila_existente is not None:
                if modo == "ignorar":
                    ignorados += 1
                    continue
                conn.execute(
                    """
                    UPDATE clientes
                    SET nombre=?, tipo_documento=?, documento=?, direccion=?,
                        codigo_postal=?, provincia=?, telefono=?, email=?, notas=?
                    WHERE id=?
                    """,
                    (
                        registro["nombre"], registro["tipo_documento"],
                        registro["documento"], registro["direccion"],
                        registro["codigo_postal"], registro["provincia"],
                        registro["telefono"], registro["email"],
                        registro["notas"], fila_existente[0],
                    ),
                )
                actualizados += 1
            else:
                conn.execute(
                    """
                    INSERT INTO clientes (nombre, tipo_documento, documento,
                        direccion, codigo_postal, provincia, telefono, email, notas)
                    VALUES (?,?,?,?,?,?,?,?,?)
                    """,
                    (
                        registro["nombre"], registro["tipo_documento"],
                        registro["documento"], registro["direccion"],
                        registro["codigo_postal"], registro["provincia"],
                        registro["telefono"], registro["email"],
                        registro["notas"],
                    ),
                )
                insertados += 1
        conn.commit()
    except sqlite3.Error as e:
        conn.rollback()
        raise BackupError(f"Error importando clientes: {e}") from e
    finally:
        conn.close()
    return _resumen("clientes", insertados, actualizados, ignorados, respaldo)


def _lanzar_errores(tipo, errores):
    """Construye el mensaje descriptivo y lanza BackupError sin cambios.

    El formato debe recordar al usuario que nada se ha modificado y, si
    hay muchas filas inválidas, mostrar las primeras y cuántas más hay.
    """
    visibles = errores[:15]
    resto = len(errores) - len(visibles)
    texto = (
        f"El fichero de {tipo} tiene {len(errores)} fila(s) inválida(s); "
        "no se importó nada. Revisa:\n\n- "
        + "\n- ".join(visibles)
    )
    if resto > 0:
        texto += f"\n…y {resto} fila(s) más con errores."
    texto += (
        "\n\nUsa la plantilla de ejemplo o corrige los datos e inténtalo "
        "de nuevo."
    )
    raise BackupError(texto)