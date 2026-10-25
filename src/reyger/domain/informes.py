"""Informes de ventas por periodo, sin interfaz gráfica.

Solo cuentan los documentos **emitidos** (tickets y facturas); los
borradores no son ventas. Los importes se suman en céntimos enteros
(``lineas_documento``) y se pasan a euros al final, así que los totales son
exactos. ``documentos.fecha`` es hora local ISO, de modo que ``date(fecha)``
es el día de la venta.
"""

from dataclasses import dataclass, field
from datetime import date
from typing import List, Tuple

from ..core import db
from .fiscal import a_euros

_DESDE = (
    "FROM lineas_documento l JOIN documentos d ON d.id = l.documento_id "
    "WHERE d.estado = 'emitido' AND date(d.fecha) BETWEEN ? AND ?"
)


@dataclass
class InformeVentas:
    desde: date
    hasta: date
    documentos: int = 0
    total: float = 0.0
    base: float = 0.0
    cuota: float = 0.0
    #: [(tipo, documentos, total)] — tipo es «ticket» o «factura».
    por_tipo: List[Tuple[str, int, float]] = field(default_factory=list)
    #: [(método de pago, documentos, total)]
    por_metodo: List[Tuple[str, int, float]] = field(default_factory=list)
    #: [(día AAAA-MM-DD, documentos, total)] en orden cronológico.
    por_dia: List[Tuple[str, int, float]] = field(default_factory=list)
    #: [(producto, unidades, total)] de más a menos facturado.
    por_producto: List[Tuple[str, int, float]] = field(default_factory=list)


def _agrupado(columna, orden, parametros):
    filas = db.query(
        f"SELECT {columna} AS clave, COUNT(DISTINCT d.id) AS docs, "
        f"SUM(l.subtotal_cent) AS total {_DESDE} GROUP BY clave ORDER BY {orden}",
        parametros,
    )
    return [(f["clave"], f["docs"], a_euros(f["total"])) for f in filas]


def informe_ventas(desde, hasta):
    """Resume las ventas emitidas entre *desde* y *hasta* (ambos incluidos)."""
    if desde > hasta:
        raise ValueError("La fecha inicial es posterior a la final")
    parametros = (desde.isoformat(), hasta.isoformat())
    totales = db.query_one(
        "SELECT COUNT(DISTINCT d.id) AS docs, SUM(l.subtotal_cent) AS total, "
        f"SUM(l.base_cent) AS base, SUM(l.cuota_cent) AS cuota {_DESDE}",
        parametros,
    )
    productos = db.query(
        "SELECT l.nombre AS nombre, SUM(l.cantidad) AS unidades, "
        f"SUM(l.subtotal_cent) AS total {_DESDE} "
        "GROUP BY l.nombre ORDER BY total DESC, l.nombre",
        parametros,
    )
    return InformeVentas(
        desde=desde,
        hasta=hasta,
        documentos=totales["docs"] or 0,
        total=a_euros(totales["total"]),
        base=a_euros(totales["base"]),
        cuota=a_euros(totales["cuota"]),
        por_tipo=_agrupado("d.tipo", "clave", parametros),
        por_metodo=_agrupado("d.metodo_pago", "total DESC", parametros),
        por_dia=_agrupado("date(d.fecha)", "clave", parametros),
        por_producto=[
            (f["nombre"], f["unidades"] or 0, a_euros(f["total"])) for f in productos
        ],
    )


def exportar_excel(informe, ruta):
    """Guarda *informe* en un libro .xlsx con una hoja por desglose."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    libro = Workbook()
    negrita = Font(bold=True)

    importe = "#,##0.00"

    def hoja(titulo, cabecera, filas, primera=False, formatos=None):
        """*formatos*: {índice de columna (1..n): formato numérico}."""
        h = libro.active if primera else libro.create_sheet()
        h.title = titulo
        h.append(cabecera)
        for celda in h[1]:
            celda.font = negrita
        for fila in filas:
            h.append(list(fila))
        for indice, formato in (formatos or {}).items():
            for (celda,) in h.iter_rows(min_row=2, min_col=indice, max_col=indice):
                if isinstance(celda.value, (int, float)):
                    celda.number_format = formato
        for columna in h.columns:
            ancho = max(len(str(c.value or "")) for c in columna)
            h.column_dimensions[columna[0].column_letter].width = min(max(ancho + 2, 12), 50)
        return h

    periodo = f"{informe.desde:%d/%m/%Y} – {informe.hasta:%d/%m/%Y}"
    resumen = hoja("Resumen", ["Concepto", "Valor"], [
        ("Periodo", periodo),
        ("Documentos emitidos", informe.documentos),
        ("Base imponible", informe.base),
        ("Cuota IVA", informe.cuota),
        ("Total", informe.total),
    ] + [(f"Total {tipo}s", total) for tipo, _, total in informe.por_tipo], primera=True)
    for (celda,) in resumen.iter_rows(min_row=4, min_col=2, max_col=2):
        celda.number_format = importe
    dias = hoja(
        "Por día", ["Día", "Documentos", "Total"],
        [(date.fromisoformat(dia), docs, total) for dia, docs, total in informe.por_dia],
        formatos={3: importe},
    )
    for (celda,) in dias.iter_rows(min_row=2, min_col=1, max_col=1):
        celda.number_format = "DD/MM/YYYY"
    hoja("Por producto", ["Producto", "Unidades", "Total"], informe.por_producto,
         formatos={3: importe})
    hoja("Por método de pago", ["Método", "Documentos", "Total"], informe.por_metodo,
         formatos={3: importe})
    libro.save(ruta)
    return ruta
